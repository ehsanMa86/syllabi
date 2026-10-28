# Copyright (c) 2026 Ehsan Madadi
# SPDX-License-Identifier: GPL-3.0-or-later
"""Run the browser smoke test in headless Chrome or Edge and check the generated .docx.

    python scripts/build.py
    python tests/run_browser_tests.py [--browser PATH] [--keep]

Needs python-docx (pip install python-docx) to inspect the Word output.
"""
from __future__ import annotations

import argparse
import base64
import functools
import html
import http.server
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
CANDIDATES = [
    "msedge", "google-chrome", "chromium", "chromium-browser", "chrome",
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
]


def find_browser(explicit: str | None) -> str:
    for c in ([explicit] if explicit else CANDIDATES):
        path = shutil.which(c) or (c if Path(c).exists() else None)
        if path:
            return path
    sys.exit("No Chrome or Edge found; pass --browser PATH")


class QuietHandler(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *args) -> None:
        pass


def serve() -> tuple[http.server.ThreadingHTTPServer, int]:
    handler = functools.partial(QuietHandler, directory=str(ROOT))
    httpd = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    return httpd, httpd.server_address[1]


def dump_dom(browser: str, url: str) -> str:
    extra = ["--no-sandbox"] if os.environ.get("CI") else []
    with tempfile.TemporaryDirectory() as profile:
        res = subprocess.run(
            [browser, "--headless=new", "--disable-gpu", "--no-first-run", f"--user-data-dir={profile}",
             *extra, "--virtual-time-budget=20000", "--dump-dom", url],
            capture_output=True, text=True, encoding="utf-8", timeout=120)
    return res.stdout


def pre(dom: str, pre_id: str) -> str:
    m = re.search(rf'<pre id="{pre_id}">(.*?)</pre>', dom, re.S)
    return html.unescape(m.group(1)) if m else ""


def check_docx(data: bytes, keep: Path | None) -> list[str]:
    import io
    import docx  # python-docx

    if keep:
        keep.write_bytes(data)
    d = docx.Document(io.BytesIO(data))
    problems = []
    styles = [p.style.name if p.style is not None else "Normal" for p in d.paragraphs]
    if "Title" not in styles:
        problems.append("no Title paragraph")
    if styles.count("Heading 1") < 4:
        problems.append(f"expected several Heading 1 paragraphs, found {styles.count('Heading 1')}")
    if not d.tables:
        problems.append("no tables")
    if not d.core_properties.title:
        problems.append("document title property is empty")
    xml = d.element.xml
    if "w:hyperlink" not in xml:
        problems.append("no hyperlinks")
    if "w:numPr" not in xml:
        problems.append("lists are not using Word numbering")
    if not any("w:tblHeader" in t._tbl.xml for t in d.tables):
        problems.append("table header rows do not repeat")
    return problems


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--browser")
    ap.add_argument("--keep", type=Path, help="save the generated .docx here")
    args = ap.parse_args()
    if not (ROOT / "dist" / "data" / "semesters.json").exists():
        sys.exit("Build first: python scripts/build.py")
    browser = find_browser(args.browser)
    httpd, port = serve()
    failures = 0
    try:
        for key in ("chhs.ccjem", "ced", "cob", "coe"):
            dom = dump_dom(browser, f"http://127.0.0.1:{port}/tests/js/smoke.html?key={key}")
            out = pre(dom, "out")
            ok = out.strip().endswith("OK")
            print(f"[{key}] {'pass' if ok else 'FAIL'}\n  " + out.replace("\n", "\n  "))
            if not ok:
                failures += 1
                continue
            b64 = pre(dom, "docx").strip()
            keep = args.keep.with_name(f"{args.keep.stem}-{key}{args.keep.suffix}") if args.keep else None
            problems = check_docx(base64.b64decode(b64), keep)
            for p in problems:
                print(f"  docx problem: {p}")
            failures += bool(problems)
    finally:
        httpd.shutdown()
    print("all browser tests passed" if not failures else f"{failures} failure(s)")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
